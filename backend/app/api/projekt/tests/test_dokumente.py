"""Stufe 5: Volltextindex, Office-Links, Vorlagen, WOPI-Host."""

import sqlite3
import time

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.deps import get_current_active_user
from app.api.projekt import wopi
from app.api.projekt.core import index, office, ordner, projekte, vorlagen, zeit


def _projekt(conn, **ueber):
    felder = dict(name="Index Test", honorarmodell="pauschal", akteur="pytest")
    felder.update(ueber)
    return projekte.anlegen(conn, **felder)


def _dateien_anlegen(o):
    (o.pfad / "01_Grundlagen" / "notiz.md").write_text("Das Bodengutachten zeigt Tonschichten.", encoding="utf-8")
    import docx
    d = docx.Document()
    d.add_paragraph("Protokoll Ortstermin: Schachtdeckel Nr. 17 ist gebrochen.")
    d.save(str(o.pfad / "03_Schriftverkehr" / "Protokoll.docx"))
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Position", "Menge"])
    wb.active.append(["Rohrleitung DN 300", 120])
    wb.save(str(o.pfad / "04_Berechnungen" / "Mengen.xlsx"))
    import fitz
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Genehmigungsbescheid Wasserrecht")
    pdf.save(str(o.pfad / "00_Vertrag" / "Bescheid.pdf"))
    (o.pfad / "00_Vertrag" / "plan.dwg").write_bytes(b"\x00" * 10)
    (o.pfad / ordner.AKTE / "geheim.txt").write_text("nicht indizieren", encoding="utf-8")


def test_index_und_suche(frische_db, app_conn, projekte_wurzel, tmp_path, monkeypatch):
    p = _projekt(app_conn)
    o = ordner.finde(p["id"])
    _dateien_anlegen(o)
    mail_db = tmp_path / "quagg.db"
    c = sqlite3.connect(mail_db)
    c.execute("CREATE TABLE email_events (id INTEGER PRIMARY KEY, subject TEXT, sender TEXT, received_at TEXT,"
              " project_id INTEGER, body_text TEXT)")
    c.execute("INSERT INTO email_events VALUES (1, 'Nachfrage Bauamt', 'amt@stadt.de', '2026-08-20 09:00:00', ?,"
              " 'Bitte den Bauzeitenplan bis Freitag schicken.')", (p["id"],))
    c.commit(); c.close()
    monkeypatch.setenv("QUAGG_SQLITE_PFAD", str(mail_db))

    b = index.aktualisieren(o, p["id"])
    assert b["neu"] == 4 and b["mails_neu"] == 1 and b["entfernt"] == 0
    assert index.stand(o)["eintraege"] == 6            # 4 Dateien + dwg (meta ohne Text) + 1 Mail
    assert [t["pfad"] for t in index.suche(o, "Tonschichten")] == ["01_Grundlagen/notiz.md"]
    assert index.suche(o, "schachtdeckel gebrochen")[0]["pfad"] == "03_Schriftverkehr/Protokoll.docx"
    assert index.suche(o, "DN 300")[0]["pfad"] == "04_Berechnungen/Mengen.xlsx"
    assert index.suche(o, "Wasserrecht")[0]["pfad"] == "00_Vertrag/Bescheid.pdf"
    mail = index.suche(o, "Bauzeitenplan")
    assert mail[0]["typ"] == "mail" and "Bauamt" in mail[0]["titel"]
    assert index.suche(o, "geheim") == [] and index.suche(o, "") == []
    assert "[" in index.suche(o, "Tonschichten")[0]["snippet"]
    # inkrementell: nichts neu; geaenderte Datei wird neu gelesen; geloeschte verschwindet
    assert index.aktualisieren(o, p["id"])["neu"] == 0
    time.sleep(0.01)
    (o.pfad / "01_Grundlagen" / "notiz.md").write_text("Jetzt Kies statt Ton.", encoding="utf-8")
    (o.pfad / "00_Vertrag" / "Bescheid.pdf").unlink()
    b = index.aktualisieren(o, p["id"], mit_mails=False)
    assert b["neu"] == 1 and b["entfernt"] == 1
    assert index.suche(o, "Tonschichten") == [] and index.suche(o, "Kies")[0]["pfad"] == "01_Grundlagen/notiz.md"
    assert index.suche(o, "Wasserrecht") == []


def test_office_links(monkeypatch):
    monkeypatch.setenv("PROJEKTE_WEBDAV_URL", "")
    assert office.office_link("01_Laufend/1338_x/00_Vertrag/Angebot.docx") is None
    assert office.konfiguration()["webdav_url"] == ""
    monkeypatch.setenv("PROJEKTE_WEBDAV_URL", "https://u1.your-storagebox.de/1_Projekte/")
    l = office.office_link("01_Laufend/1338_x/00_Vertrag/Angebot Nr 1.docx")
    assert l["schema"] == "ms-word"
    assert l["uri"] == "ms-word:ofe|u|https://u1.your-storagebox.de/1_Projekte/01_Laufend/1338_x/00_Vertrag/Angebot%20Nr%201.docx"
    assert office.office_link("x/y.xlsx")["schema"] == "ms-excel"
    assert office.office_link("x/y.dwg") is None


def test_vorlagen_erzeugen(frische_db, app_conn, projekte_wurzel, tmp_path, monkeypatch):
    monkeypatch.setattr(vorlagen, "VORLAGEN_ORDNER", tmp_path / "vorlagen")
    p = _projekt(app_conn, name="Vorlagen Test", kurzname="VT")
    projekte.beteiligter_anlegen(app_conn, p["id"], rolle="auftraggeber", name="Stadt Musterhausen",
                                 kontakt="Rathausplatz 1", akteur="pytest")
    zeit.buchen(app_conn, p["id"], datum=__import__("datetime").date(2027, 2, 1), dauer_min=90,
                taetigkeit="Vermessung", akteur="pytest")
    akte = projekte.lesen(app_conn, p["id"])
    firma = {"name": "Quagg Engineering", "strasse": "Musterstr. 1", "plz": "56070", "ort": "Koblenz", "email": "x@y.de"}
    o = ordner.finde(p["id"])

    a = vorlagen.erzeugen(app_conn, akte, "anschreiben", akteur="fabio", firma=firma)
    assert a["pfad"].startswith("03_Schriftverkehr/Anschreiben_") and a["pfad"].endswith(".docx")
    import docx
    text = "\n".join(par.text for par in docx.Document(str(o.pfad / a["pfad"])).paragraphs)
    assert "Stadt Musterhausen" in text and "Quagg Engineering" in text and "#P" in text and "{{" not in text
    b = vorlagen.erzeugen(app_conn, akte, "anschreiben", akteur="fabio", firma=firma)
    assert b["pfad"] != a["pfad"]                       # nie ueberschreiben
    n = vorlagen.erzeugen(app_conn, akte, "aktennotiz", akteur="fabio", firma=None, name="Ortstermin 12.3.")
    assert n["pfad"] == "03_Schriftverkehr/Ortstermin_12_3.docx"
    s = vorlagen.erzeugen(app_conn, akte, "stundennachweis", akteur="fabio")
    import openpyxl
    ws = openpyxl.load_workbook(str(o.pfad / s["pfad"])).active
    zeilen = list(ws.iter_rows(values_only=True))
    assert zeilen[3][0] == "Datum" and zeilen[4][1] == "Vermessung" and zeilen[-1][3] == 90
    h = vorlagen.erzeugen(app_conn, akte, "honorarermittlung", akteur="fabio")
    assert h["pfad"].startswith("00_Vertrag/Honorarermittlung_")
    with pytest.raises(vorlagen.VorlageAbgelehnt):
        vorlagen.erzeugen(app_conn, akte, "brief", akteur="fabio")
    assert {v["id"] for v in vorlagen.liste()} == {"anschreiben", "aktennotiz", "stundennachweis", "honorarermittlung"}


def _wopi_client():
    app = FastAPI()
    app.include_router(wopi.router, prefix="/FastAPI/wopi")
    return TestClient(app)


def test_wopi_host(frische_db, app_conn, projekte_wurzel, monkeypatch):
    p = _projekt(app_conn, name="WOPI")
    o = ordner.finde(p["id"])
    import docx
    d = docx.Document(); d.add_paragraph("Version 1"); d.save(str(o.pfad / "00_Vertrag" / "Auftrag.docx"))
    monkeypatch.setenv("ONLYOFFICE_URL", "")
    with pytest.raises(wopi.WopiAbgelehnt):
        wopi.session(p["id"], "00_Vertrag/Auftrag.docx", "fabio", "https://api")
    monkeypatch.setenv("ONLYOFFICE_URL", "https://office.example.org/")
    s = wopi.session(p["id"], "00_Vertrag/Auftrag.docx", "fabio", "https://api")
    assert s["editor_url"].startswith("https://office.example.org/hosting/wopi/word/edit?wopisrc=https://api/FastAPI/wopi/files/")
    fid, tok = s["file_id"], s["access_token"]
    c = _wopi_client()
    info = c.get(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok})
    assert info.status_code == 200 and info.json()["BaseFileName"] == "Auftrag.docx" and info.json()["UserCanWrite"]
    assert c.get(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok + "x"}).status_code == 401
    fremd = wopi.token_erzeugen(p["id"], "00_Vertrag/anders.docx", "fabio")
    assert c.get(f"/FastAPI/wopi/files/{fid}", params={"access_token": fremd}).status_code == 401
    inhalt = c.get(f"/FastAPI/wopi/files/{fid}/contents", params={"access_token": tok})
    assert inhalt.status_code == 200 and inhalt.content[:2] == b"PK"
    # Lock-Protokoll
    assert c.post(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok},
                  headers={"X-WOPI-Override": "LOCK", "X-WOPI-Lock": "L1"}).status_code == 200
    r = c.post(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok},
               headers={"X-WOPI-Override": "LOCK", "X-WOPI-Lock": "L2"})
    assert r.status_code == 409 and r.headers["X-WOPI-Lock"] == "L1"
    assert c.post(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok},
                  headers={"X-WOPI-Override": "GET_LOCK"}).headers["X-WOPI-Lock"] == "L1"
    # Speichern mit falschem Lock -> 409, mit richtigem -> Datei ersetzt
    neu = docx.Document(); neu.add_paragraph("Version 2")
    import io
    puffer = io.BytesIO(); neu.save(puffer)
    assert c.post(f"/FastAPI/wopi/files/{fid}/contents", params={"access_token": tok},
                  headers={"X-WOPI-Override": "PUT", "X-WOPI-Lock": "L2"}, content=puffer.getvalue()).status_code == 409
    ok = c.post(f"/FastAPI/wopi/files/{fid}/contents", params={"access_token": tok},
                headers={"X-WOPI-Override": "PUT", "X-WOPI-Lock": "L1"}, content=puffer.getvalue())
    assert ok.status_code == 200
    assert docx.Document(str(o.pfad / "00_Vertrag" / "Auftrag.docx")).paragraphs[0].text == "Version 2"
    assert c.post(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok},
                  headers={"X-WOPI-Override": "UNLOCK", "X-WOPI-Lock": "L1"}).status_code == 200
    assert c.post(f"/FastAPI/wopi/files/{fid}", params={"access_token": tok},
                  headers={"X-WOPI-Override": "GET_LOCK"}).headers["X-WOPI-Lock"] == ""
    # Pfad-Schutz
    boese = wopi.file_id(p["id"], "../../x.docx")
    assert c.get(f"/FastAPI/wopi/files/{boese}", params={"access_token": wopi.token_erzeugen(p['id'], '../../x.docx', 'f')}).status_code == 403
    akte_fid = wopi.file_id(p["id"], "_akte/akte.docx")
    assert c.get(f"/FastAPI/wopi/files/{akte_fid}", params={"access_token": wopi.token_erzeugen(p['id'], '_akte/akte.docx', 'f')}).status_code == 403


def test_router_dokumente(frische_db, projekte_wurzel, monkeypatch):
    from app.api.projekt.router import router
    app = FastAPI()
    app.include_router(router, prefix="/FastAPI/projekte")
    app.dependency_overrides[get_current_active_user] = lambda: __import__("types").SimpleNamespace(role="INTERNAL", username="pytest")
    c = TestClient(app)
    monkeypatch.setenv("PROJEKTE_WEBDAV_URL", "https://u1.your-storagebox.de/1_Projekte")
    monkeypatch.setenv("ONLYOFFICE_URL", "")
    pid = c.post("/FastAPI/projekte", json={"name": "Dok", "honorarmodell": "pauschal"}).json()["id"]
    assert c.get("/FastAPI/projekte/konfiguration").json()["webdav_url"].endswith("1_Projekte")
    assert len(c.get("/FastAPI/projekte/vorlagen").json()) == 4
    o = ordner.finde(pid)
    (o.pfad / "01_Grundlagen" / "hinweis.txt").write_text("Regenrückhaltebecken Volumen 800 m³", encoding="utf-8")
    assert c.post(f"/FastAPI/projekte/{pid}/index").json()["neu"] == 1
    t = c.get(f"/FastAPI/projekte/{pid}/suche", params={"q": "Volumen"}).json()
    assert t["treffer"][0]["pfad"] == "01_Grundlagen/hinweis.txt" and t["index"]["eintraege"] == 1
    assert c.get(f"/FastAPI/projekte/{pid}/suche", params={"q": "V"}).status_code == 422
    l = c.get(f"/FastAPI/projekte/{pid}/office-link", params={"pfad": "00_Vertrag/Angebot.docx"}).json()
    assert l["uri"].startswith("ms-word:ofe|u|https://u1.your-storagebox.de/1_Projekte/00_Angebote/")
    assert c.get(f"/FastAPI/projekte/{pid}/office-link", params={"pfad": "../x.docx"}).status_code == 422
    monkeypatch.setattr(vorlagen, "VORLAGEN_ORDNER", projekte_wurzel.parent / "vorlagen")
    v = c.post(f"/FastAPI/projekte/{pid}/vorlagen/aktennotiz", json={"name": "Notiz"})
    assert v.status_code == 201 and v.json()["pfad"] == "03_Schriftverkehr/Notiz.docx"
    assert c.post(f"/FastAPI/projekte/{pid}/vorlagen/quatsch", json={}).status_code == 422
    assert c.get(f"/FastAPI/projekte/{pid}/wopi-session", params={"pfad": "03_Schriftverkehr/Notiz.docx"}).status_code == 503
