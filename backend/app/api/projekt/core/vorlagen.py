"""Vorlagen-Generator: docx (docxtpl/Jinja) und xlsx (openpyxl) mit Projektdaten,
geschrieben in den Projektordner — danach im Desktop-Office weiterbearbeiten.

Die docx-Vorlagen werden beim ersten Gebrauch mit python-docx erzeugt (keine
Binaerdateien im Git); wer eigene Vorlagen will, ersetzt die Dateien unter
backend/app/api/projekt/vorlagen/ (Jinja-Felder siehe _KONTEXT_HINWEIS).
"""

import re
from datetime import date
from pathlib import Path

from . import ordner, zeit

VORLAGEN_ORDNER = Path(__file__).resolve().parents[1] / "vorlagen"

VORLAGEN = {
    "anschreiben": {"titel": "Anschreiben (Word)", "typ": "docx", "zielordner": "03_Schriftverkehr",
                    "datei": "Anschreiben.docx"},
    "aktennotiz": {"titel": "Aktennotiz (Word)", "typ": "docx", "zielordner": "03_Schriftverkehr",
                   "datei": "Aktennotiz.docx"},
    "stundennachweis": {"titel": "Stundennachweis (Excel)", "typ": "xlsx", "zielordner": "05_Rechnungen",
                        "datei": "Stundennachweis.xlsx"},
    "honorarermittlung": {"titel": "Honorarermittlung (Excel)", "typ": "xlsx", "zielordner": "00_Vertrag",
                          "datei": "Honorarermittlung.xlsx"},
}

_KONTEXT_HINWEIS = ("Felder: {{ projekt.id }}, {{ projekt.name }}, {{ projekt.kurzname }}, {{ datum }},"
                    " {{ firma.name }}, {{ firma.strasse }}, {{ firma.plz }} {{ firma.ort }}, {{ firma.email }},"
                    " {{ auftraggeber.name }}, {{ auftraggeber.kontakt }}, {{ bauherr.name }}, {{ bearbeiter }}")


class VorlageAbgelehnt(ValueError):
    """Router: 422."""


def liste() -> list[dict]:
    return [{"id": k, **v} for k, v in VORLAGEN.items()]


# ── docx-Vorlagen erzeugen (einmalig) ────────────────────────────────────────

def _docx_vorlage_schreiben(ziel: Path, art: str) -> None:
    import docx
    d = docx.Document()
    d.add_paragraph("{{ firma.name }} · {{ firma.strasse }} · {{ firma.plz }} {{ firma.ort }}")
    d.add_paragraph("")
    if art == "anschreiben":
        d.add_paragraph("{{ auftraggeber.name }}")
        d.add_paragraph("{{ auftraggeber.kontakt }}")
        d.add_paragraph("")
        d.add_paragraph("{{ firma.ort }}, {{ datum }}")
        d.add_paragraph("")
        d.add_paragraph("Projekt #P{{ projekt.id }} {{ projekt.name }}").runs[0].bold = True
        d.add_paragraph("")
        d.add_paragraph("Sehr geehrte Damen und Herren,")
        d.add_paragraph("")
        d.add_paragraph("…")
        d.add_paragraph("")
        d.add_paragraph("Mit freundlichen Grüßen")
        d.add_paragraph("{{ bearbeiter }}")
    else:
        d.add_paragraph("Aktennotiz").runs[0].bold = True
        d.add_paragraph("Projekt: #P{{ projekt.id }} {{ projekt.name }}")
        d.add_paragraph("Datum: {{ datum }} · Verfasser: {{ bearbeiter }}")
        d.add_paragraph("Beteiligte: {{ bauherr.name }} / {{ auftraggeber.name }}")
        d.add_paragraph("")
        d.add_paragraph("Anlass:")
        d.add_paragraph("…")
        d.add_paragraph("")
        d.add_paragraph("Ergebnis / Vereinbarung:")
        d.add_paragraph("…")
    d.add_paragraph("")
    d.add_paragraph(_KONTEXT_HINWEIS).runs[0].font.size = docx.shared.Pt(7)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(ziel))


def sicherstellen() -> None:
    for k, v in VORLAGEN.items():
        if v["typ"] == "docx" and not (VORLAGEN_ORDNER / v["datei"]).is_file():
            _docx_vorlage_schreiben(VORLAGEN_ORDNER / v["datei"], k)


# ── Kontext ──────────────────────────────────────────────────────────────────

def _beteiligter(projekt: dict, rolle: str) -> dict:
    for b in projekt.get("beteiligte", []):
        if b["rolle"] == rolle:
            return {"name": b["name"], "kontakt": b.get("kontakt", "")}
    return {"name": "", "kontakt": ""}


def kontext(projekt: dict, firma: dict | None, bearbeiter: str) -> dict:
    firma = firma or {}
    return {
        "projekt": {"id": projekt["id"], "name": projekt["name"], "kurzname": projekt.get("kurzname", ""),
                    "phase": projekt.get("phase") or ""},
        "firma": {k: str(firma.get(k, "") or "") for k in ("name", "strasse", "plz", "ort", "email", "telefon", "ust_id", "iban")},
        "auftraggeber": _beteiligter(projekt, "auftraggeber") if _beteiligter(projekt, "auftraggeber")["name"]
        else _beteiligter(projekt, "bauherr"),
        "bauherr": _beteiligter(projekt, "bauherr"),
        "bearbeiter": bearbeiter,
        "datum": date.today().strftime("%d.%m.%Y"),
    }


# ── Erzeugen ─────────────────────────────────────────────────────────────────

def _freier_name(ordnerpfad: Path, name: str, endung: str) -> Path:
    ziel = ordnerpfad / f"{name}{endung}"
    n = 2
    while ziel.exists():
        ziel = ordnerpfad / f"{name}-{n}{endung}"
        n += 1
    return ziel


def _sicher(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", text).strip("_")[:40] or "Dokument"


def _xlsx_stundennachweis(conn, projekt: dict, ziel: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stundennachweis"
    ws.append([f"Stundennachweis #P{projekt['id']} {projekt['name']}"])
    ws.append([f"Stand {date.today().strftime('%d.%m.%Y')}"])
    ws.append([])
    ws.append(["Datum", "Tätigkeit", "Abschnitt", "Minuten", "Stunden", "abrechenbar", "Rechnung"])
    namen = {a["id"]: a["bezeichnung"] for a in projekt.get("abschnitte", [])}
    buchungen = sorted(zeit.liste(conn, projekt["id"]), key=lambda z: (z["datum"], z["id"]))
    for z in buchungen:
        ws.append([z["datum"], z["taetigkeit"], namen.get(z["abschnitt_id"], ""), z["dauer_min"],
                   round(z["dauer_min"] / 60, 2), "ja" if z["abrechenbar"] else "nein",
                   f"#{z['rechnung_id']}" if z["rechnung_id"] else ""])
    summe = sum(z["dauer_min"] for z in buchungen)
    ws.append(["Summe", "", "", summe, round(summe / 60, 2), "", ""])
    for spalte, breite in zip("ABCDEFG", (12, 44, 28, 10, 10, 12, 10)):
        ws.column_dimensions[spalte].width = breite
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=13)
    for zelle in ws[4]:
        zelle.font = openpyxl.styles.Font(bold=True)
    wb.save(str(ziel))


def _xlsx_honorarermittlung(projekt: dict, ziel: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Honorar"
    ws.append([f"Honorarermittlung #P{projekt['id']} {projekt['name']}"])
    ws.append([f"Honorarmodell: {projekt.get('honorarmodell')}"
               + (f" · Leistungsbild § {projekt['leistungsbild']}" if projekt.get("leistungsbild") else "")])
    ws.append([])
    ws.append(["Nr", "Abschnitt", "Art", "LPH", "Honorar netto €", "beauftragt", "Fortschritt %", "Leistung netto €", "Status"])
    start = ws.max_row + 1
    for a in projekt.get("abschnitte", []):
        zeile = ws.max_row + 1
        ws.append([a["nr"], a["bezeichnung"], a["art"], a["lph"] or "", a["honorar_cent"] / 100,
                   "ja" if a["beauftragt"] else "nein", a["fortschritt_prozent"],
                   f"=E{zeile}*G{zeile}/100", a["status"]])
    ende = ws.max_row
    if ende >= start:
        ws.append(["", "Summe", "", "", f"=SUM(E{start}:E{ende})", "", "", f"=SUM(H{start}:H{ende})", ""])
    for spalte, breite in zip("ABCDEFGHI", (5, 40, 12, 6, 16, 11, 13, 16, 12)):
        ws.column_dimensions[spalte].width = breite
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=13)
    for zelle in ws[4]:
        zelle.font = openpyxl.styles.Font(bold=True)
    for zeile in ws.iter_rows(min_row=start, max_row=ws.max_row, min_col=5, max_col=8):
        for zelle in zeile:
            if zelle.column in (5, 8):
                zelle.number_format = "#,##0.00"
    wb.save(str(ziel))


def erzeugen(conn, projekt: dict, vorlage_id: str, *, akteur: str, firma: dict | None = None,
             name: str | None = None) -> dict:
    v = VORLAGEN.get(vorlage_id)
    if v is None:
        raise VorlageAbgelehnt(f"unbekannte vorlage {vorlage_id!r} (bekannt: {', '.join(VORLAGEN)})")
    o = ordner.finde(projekt["id"])
    if o is None:
        raise VorlageAbgelehnt("projekt hat keinen ordner")
    zielordner = o.pfad / v["zielordner"]
    zielordner.mkdir(exist_ok=True)
    basisname = _sicher(name) if name else f"{v['datei'].rsplit('.', 1)[0]}_{date.today().isoformat()}"
    ziel = _freier_name(zielordner, basisname, "." + v["typ"])
    if v["typ"] == "docx":
        sicherstellen()
        from docxtpl import DocxTemplate
        doc = DocxTemplate(str(VORLAGEN_ORDNER / v["datei"]))
        doc.render(kontext(projekt, firma, akteur))
        doc.save(str(ziel))
    elif vorlage_id == "stundennachweis":
        _xlsx_stundennachweis(conn, projekt, ziel)
    else:
        _xlsx_honorarermittlung(projekt, ziel)
    rel = str(ziel.relative_to(o.pfad))
    return {"vorlage": vorlage_id, "pfad": rel, "projektpfad": f"{o.phase}/{o.ordnername}/{rel}",
            "groesse": ziel.stat().st_size}
