"""Volltextindex je Projekt: <Projekt>/_akte/index.sqlite (FTS5), Muster ingest/.

Inkrementell ueber mtime; pdf (pymupdf), docx (python-docx), xlsx (openpyxl),
Textdateien direkt, E-Mails aus email_events. Groessen- und Anzahl-Deckel,
damit ein Riesenordner den Aufruf nicht minutenlang haelt.
"""

import os
import sqlite3
import time
from pathlib import Path

from . import mails, ordner

DATEINAME = "index.sqlite"
TEXT_ENDUNGEN = {".txt", ".md", ".csv", ".yaml", ".yml", ".json", ".log"}
MAX_DATEI_BYTES = 50 * 1024 * 1024
MAX_ZEICHEN = 400_000
MAX_DATEIEN = 3000


def _pdf(pfad: Path) -> str:
    import fitz
    teile = []
    with fitz.open(pfad) as doc:
        for seite in doc:
            teile.append(seite.get_text())
            if sum(len(t) for t in teile) > MAX_ZEICHEN:
                break
    return "\n".join(teile)


def _docx(pfad: Path) -> str:
    import docx
    d = docx.Document(str(pfad))
    teile = [p.text for p in d.paragraphs]
    for tabelle in d.tables:
        for zeile in tabelle.rows:
            teile.append(" | ".join(z.text for z in zeile.cells))
    return "\n".join(teile)


def _xlsx(pfad: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(pfad), read_only=True, data_only=True)
    teile = []
    for ws in wb.worksheets:
        teile.append(f"# {ws.title}")
        for zeile in ws.iter_rows(values_only=True):
            werte = [str(v) for v in zeile if v is not None and str(v).strip()]
            if werte:
                teile.append(" | ".join(werte))
            if sum(len(t) for t in teile) > MAX_ZEICHEN:
                break
    wb.close()
    return "\n".join(teile)


def text_aus(pfad: Path) -> str | None:
    """None = Format wird nicht indiziert (dwg, Bilder, …)."""
    endung = pfad.suffix.lower()
    try:
        if endung == ".pdf":
            text = _pdf(pfad)
        elif endung == ".docx":
            text = _docx(pfad)
        elif endung == ".xlsx":
            text = _xlsx(pfad)
        elif endung in TEXT_ENDUNGEN:
            text = pfad.read_text(encoding="utf-8", errors="replace")
        else:
            return None
    except Exception as fehler:  # noqa: BLE001 — eine kaputte Datei stoppt nicht den Index
        return f"[nicht lesbar: {type(fehler).__name__}: {fehler}]"
    return text[:MAX_ZEICHEN]


def _verbindung(o: ordner.Ordner) -> sqlite3.Connection:
    pfad = ordner.akte_sicherstellen(o) / DATEINAME
    conn = sqlite3.connect(pfad)
    conn.execute("CREATE TABLE IF NOT EXISTS meta (pfad TEXT PRIMARY KEY, mtime REAL NOT NULL,"
                 " groesse INTEGER NOT NULL, indiziert_am REAL NOT NULL)")
    conn.execute("CREATE VIRTUAL TABLE IF NOT EXISTS fts USING fts5(pfad UNINDEXED, titel, inhalt,"
                 " tokenize='unicode61 remove_diacritics 2')")
    return conn


def _dateien(o: ordner.Ordner):
    zaehler = 0
    for wurzel, dirs, files in os.walk(o.pfad):
        dirs[:] = sorted(d for d in dirs if not d.startswith(".") and d != ordner.AKTE)
        for f in sorted(files):
            if f.startswith(".") or f.startswith("~$"):
                continue
            p = Path(wurzel) / f
            try:
                st = p.stat()
            except OSError:
                continue
            zaehler += 1
            if zaehler > MAX_DATEIEN:
                return
            yield str(p.relative_to(o.pfad)), p, st


def aktualisieren(o: ordner.Ordner, projekt_id: int, *, mit_mails: bool = True) -> dict:
    start = time.monotonic()
    conn = _verbindung(o)
    try:
        bekannt = {pfad: (mtime, groesse) for pfad, mtime, groesse in conn.execute("SELECT pfad, mtime, groesse FROM meta")}
        gesehen: set[str] = set()
        neu = uebersprungen = 0
        for rel, p, st in _dateien(o):
            gesehen.add(rel)
            if st.st_size > MAX_DATEI_BYTES:
                uebersprungen += 1
                continue
            if bekannt.get(rel) == (st.st_mtime, st.st_size):
                continue
            text = text_aus(p)
            conn.execute("DELETE FROM fts WHERE pfad = ?", (rel,))
            if text is not None:
                conn.execute("INSERT INTO fts (pfad, titel, inhalt) VALUES (?, ?, ?)", (rel, p.name, text))
                neu += 1
            conn.execute("INSERT OR REPLACE INTO meta (pfad, mtime, groesse, indiziert_am) VALUES (?, ?, ?, ?)",
                         (rel, st.st_mtime, st.st_size, time.time()))
        entfernt = 0
        for rel in list(bekannt):
            if rel.startswith("mail:") or rel in gesehen:
                continue
            conn.execute("DELETE FROM fts WHERE pfad = ?", (rel,))
            conn.execute("DELETE FROM meta WHERE pfad = ?", (rel,))
            entfernt += 1
        mails_neu = 0
        if mit_mails:
            for m in mails.inhalte(projekt_id):
                rel = f"mail:{m['id']}"
                if rel in bekannt:
                    continue
                conn.execute("INSERT INTO fts (pfad, titel, inhalt) VALUES (?, ?, ?)",
                             (rel, f"E-Mail {m['empfangen_am']} {m['absender']}: {m['betreff']}", m["text"]))
                conn.execute("INSERT OR REPLACE INTO meta (pfad, mtime, groesse, indiziert_am) VALUES (?, 0, ?, ?)",
                             (rel, len(m["text"]), time.time()))
                mails_neu += 1
        conn.commit()
        anzahl = conn.execute("SELECT count(*) FROM meta").fetchone()[0]
    finally:
        conn.close()
    return {"neu": neu, "entfernt": entfernt, "mails_neu": mails_neu, "uebersprungen": uebersprungen,
            "eintraege": anzahl, "dauer_s": round(time.monotonic() - start, 2)}


def stand(o: ordner.Ordner) -> dict | None:
    pfad = o.pfad / ordner.AKTE / DATEINAME
    if not pfad.is_file():
        return None
    conn = sqlite3.connect(f"file:{pfad}?mode=ro", uri=True)
    try:
        anzahl, letzte = conn.execute("SELECT count(*), max(indiziert_am) FROM meta").fetchone()
    finally:
        conn.close()
    return {"eintraege": anzahl, "zuletzt": letzte}


def _fts_anfrage(query: str) -> str:
    """Nutzertext -> sichere FTS5-Anfrage: Woerter mit Praefix-Suche, UND-verknuepft."""
    woerter = [w.strip('"*() ') for w in query.replace("'", " ").split()]
    woerter = [w for w in woerter if w]
    if not woerter:
        return ""
    return " AND ".join(f'"{w}"*' for w in woerter)


def suche(o: ordner.Ordner, query: str, limit: int = 20) -> list[dict]:
    pfad = o.pfad / ordner.AKTE / DATEINAME
    anfrage = _fts_anfrage(query)
    if not pfad.is_file() or not anfrage:
        return []
    conn = sqlite3.connect(f"file:{pfad}?mode=ro", uri=True)
    try:
        zeilen = conn.execute(
            "SELECT pfad, titel, snippet(fts, 2, '[', ']', '…', 14), bm25(fts) FROM fts"
            " WHERE fts MATCH ? ORDER BY bm25(fts) LIMIT ?", (anfrage, limit)).fetchall()
    except sqlite3.OperationalError:
        return []
    finally:
        conn.close()
    return [{"pfad": p, "titel": t, "snippet": s, "rang": round(-r, 3), "typ": "mail" if p.startswith("mail:") else "datei"}
            for p, t, s, r in zeilen]
